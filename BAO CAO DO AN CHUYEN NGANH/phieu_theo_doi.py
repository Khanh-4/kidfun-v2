"""Điền mẫu "[2] Phieu theo doi tien do DACN.doc" -> PhieuTheoDoiTuanN.pdf.

Mỗi tuần thêm 1 mục vào WEEKS rồi chạy:  python3 phieu_theo_doi.py N
Phiếu tuần N chứa các dòng tuần 1..N. Cần Word trên Windows (WSL gọi qua PowerShell)
để chuyển .doc <-> .docx/.pdf.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "[2] Phieu theo doi tien do DACN.doc"

TEN_DE_TAI = (
    "Ứng dụng học máy dự đoán và phát hiện bất thường trong hành vi sử dụng "
    "thiết bị di động của trẻ em, tích hợp vào hệ thống giám sát KidFun"
)
GVHD = "Dương Thành Phết"
THANH_VIEN = [  # (họ tên, MSSV, lớp) — mẫu có tối đa 3 dòng
    ("Cao Duy Quốc Khánh", "2380601019", "23DTHC1"),
    ("Đinh Bùi Tuấn Anh", "2380600035", "23DTHC1"),
]
NGANH = "CÔNG NGHỆ THÔNG TIN"
CHUYEN_NGANH = "CÔNG NGHỆ PHẦN MỀM"

WEEKS = {
    1: (
        "12/10/2026\n–\n18/10/2026",
        [
            "Chốt đề tài, 4 module (trọng tâm M1 dự đoán, M2 phát hiện bất thường).",
            "Cập nhật mô tả đề tài theo 4 góp ý GVHD (cảnh báo trên máy, giải thích lý do, "
            "tiết kiệm pin, phụ huynh xác nhận đúng/sai).",
            "Cài đặt bộ thu thập dữ liệu theo giờ/theo app: bảng AppUsageHourly, API backend, "
            "app Android gửi theo lô (WorkManager); đã triển khai.",
            "EDA 2 dataset công khai LSApp, StudentLife; baseline MA7 tốt nhất.",
            "Đặc tả bộ đặc trưng M1/M2, quy trình đo pin.",
            "Tồn đọng: chờ test máy thật, số đo pin.",
        ],
    ),
}

CONTENT_PT = 11


def win(p: Path) -> str:
    return subprocess.check_output(["wslpath", "-w", str(p)], text=True).strip()


def word_convert(src: Path, dst: Path, fmt: int) -> None:
    """fmt: 16 = docx, 17 = pdf."""
    script = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{win(src)}'); $d.SaveAs2('{win(dst)}', {fmt}); "
        "$d.Close(); try { $w.Quit() } catch {}"  # Quit đôi khi báo lỗi RPC dù đã lưu xong
    )
    subprocess.run(["powershell.exe", "-NoProfile", "-Command", script], check=False)
    if not dst.exists():
        sys.exit(f"Word không tạo được {dst}")


def set_lines(cell, lines, size=CONTENT_PT):
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)
    first = paras[0]
    for i, line in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        p.add_run(line).font.size = Pt(size)


def remove(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def fill_member(p, name, mssv, lop) -> None:
    """Dòng mẫu: '(n)' tự đánh số + tab kẻ chấm; đổi sang tab trơn rồi ghi chữ."""
    tabs = p._p.get_or_add_pPr().find(qn("w:tabs"))
    if tabs is not None:
        tabs.getparent().remove(tabs)
    stops = p.paragraph_format.tab_stops
    for pos in (Twips(1100), Twips(4536), Twips(7300)):  # sau "(n)" / MSSV / Lớp
        stops.add_tab_stop(pos)
    for r in p.runs:
        r._element.getparent().remove(r._element)
    p.add_run(f"{name}\tMSSV: {mssv}\tLớp: {lop}")


def fill(doc, upto: int) -> None:
    paras = doc.paragraphs
    paras[6].runs[-1].text = TEN_DE_TAI  # thay tab kẻ chấm bằng nội dung
    remove(paras[7])  # dòng kẻ chấm thứ 2 của tên đề tài
    paras[8].runs[-1].text = GVHD
    member_lines = paras[10:13]
    for p, member in zip(member_lines, THANH_VIEN):
        fill_member(p, *member)
    for p in member_lines[len(THANH_VIEN):]:
        remove(p)
    paras[13].runs[-1].text = NGANH
    paras[14].runs[-1].text = CHUYEN_NGANH

    rows = doc.tables[0].rows
    for week in range(1, upto + 1):
        date, bullets = WEEKS[week]
        row = rows[week]
        set_lines(row.cells[1], date.split("\n"))
        set_lines(row.cells[2], [f"- {b}" for b in bullets])


def main() -> None:
    week = int(sys.argv[1]) if len(sys.argv) > 1 else max(WEEKS)
    if week not in WEEKS:
        sys.exit(f"Chưa có nội dung tuần {week} trong WEEKS")
    with tempfile.TemporaryDirectory(dir=HERE) as tmp:
        tmp = Path(tmp)
        src_docx = tmp / "mau.docx"
        word_convert(TEMPLATE, src_docx, 16)
        doc = docx.Document(src_docx)
        fill(doc, week)
        filled = tmp / "filled.docx"
        doc.save(filled)
        word_convert(filled, HERE / f"PhieuTheoDoiTuan{week}.pdf", 17)
    print(f"Đã tạo PhieuTheoDoiTuan{week}.pdf")


if __name__ == "__main__":
    main()
